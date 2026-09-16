"""Summarise the PDGui control and finite-q polariton calculations."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


LEGEND_RE = re.compile(
    r"Fig(?P<figure>9|10|11a|11b) [YZ]=(?P<offset>[0-9.]+)mm "
    r"(?P<configuration>x\([a-z]+\)x)"
)


def load_pdgui_spectra(workbook: Path):
    """Return frequency grid and scenario spectra from the PDGui workbook."""
    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Crystal Raman"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    legends = [str(value) for value in header[2:] if value is not None]
    data = [row for row in rows if row[1] is not None]
    frequencies = np.asarray([row[1] for row in data], dtype=float)
    spectra = {
        legend: np.asarray([row[index + 2] for row in data], dtype=float)
        for index, legend in enumerate(legends)
    }
    return frequencies, spectra


def summarise_pdgui(frequencies, spectra):
    """Extract the strongest q->0 NAC feature in the polar-mode interval."""
    polar_mask = (frequencies >= 180.0) & (frequencies <= 760.0)
    x = frequencies[polar_mask]
    rows = []
    for legend, spectrum in spectra.items():
        match = LEGEND_RE.fullmatch(legend)
        if match is None:
            raise ValueError(f"Unexpected PDGui legend {legend!r}")
        y = spectrum[polar_mask]
        peak_index = int(np.argmax(y))
        rows.append(
            {
                "figure": match.group("figure"),
                "configuration": match.group("configuration"),
                "offset_mm": float(match.group("offset")),
                "nac_peak_frequency_cm-1": float(x[peak_index]),
                "nac_peak_intensity": float(y[peak_index]),
                "integrated_intensity_180_760": float(np.trapezoid(y, x)),
            }
        )
    return rows


def read_csv(path: Path):
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(rows, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_results(pdgui_rows, polariton_rows, output: Path):
    """Write a concise, generated numerical report."""
    finite = [row for row in polariton_rows if row["model"] == "dft"]
    by_scan = {}
    for row in finite:
        key = (row["figure"], row["configuration"])
        by_scan.setdefault(key, []).append(float(row["peak_frequency_cm-1"]))
    nac_by_scan = {}
    for row in pdgui_rows:
        key = (row["figure"], row["configuration"])
        nac_by_scan.setdefault(key, []).append(float(row["nac_peak_frequency_cm-1"]))

    lines = [
        "# Generated verification results",
        "",
        "The finite-q values below are maxima of a projected Maxwell response,",
        "not absolute Raman intensities. The PDGui values are the strongest",
        "features of the directional q->0 `modal_pairs` spectrum.",
        "",
        "| figure | configuration | finite-q DFT peak range / cm^-1 | PDGui NAC peak range / cm^-1 |",
        "|---|---|---:|---:|",
    ]
    for key in sorted(by_scan):
        finite_values = by_scan[key]
        nac_values = nac_by_scan[key]
        lines.append(
            f"| {key[0]} | `{key[1]}` | {min(finite_values):.1f}--{max(finite_values):.1f} "
            f"| {min(nac_values):.1f}--{max(nac_values):.1f} |"
        )
    lines.extend(
        [
            "",
            "A varying finite-q peak accompanied by a nearly fixed NAC parent-mode",
            "feature is the intended validation signature. Agreement of the paper and",
            "DFT dielectric models is a numerical/implementation check; disagreement",
            "with an experimental band can also reflect the material parameters, Raman",
            "source tensor, and finite aperture averaging, none of which should be hidden",
            "inside an arbitrary post-hoc broadening.",
            "",
        ]
    )
    output.write_text("\n".join(lines))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("Data/irmer_pdgui.xlsx"))
    parser.add_argument(
        "--polariton-peaks", type=Path, default=Path("Data/polariton_peaks.csv")
    )
    parser.add_argument("--csv", type=Path, default=Path("Data/pdgui_peaks.csv"))
    parser.add_argument("--markdown", type=Path, default=Path("RESULTS.md"))
    args = parser.parse_args()

    frequencies, spectra = load_pdgui_spectra(args.workbook)
    pdgui_rows = summarise_pdgui(frequencies, spectra)
    polariton_rows = read_csv(args.polariton_peaks)
    write_csv(pdgui_rows, args.csv)
    write_results(pdgui_rows, polariton_rows, args.markdown)
    print(f"Wrote {args.csv} and {args.markdown}")


if __name__ == "__main__":
    main()
