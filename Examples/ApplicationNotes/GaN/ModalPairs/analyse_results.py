"""Summarise the PDGui GaN modal-pairs workbook without pandas."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


LEGEND_RE = re.compile(
    r"theta_ext=(?P<angle>[0-9.]+) "
    r"(?P<nac>geometry|dominant|modal_pairs) "
    r"(?P<channel>[eo]-[eo])"
)


def load_spectra(workbook: Path) -> tuple[np.ndarray, dict[str, np.ndarray]]:
    """Return the frequency grid and named spectra from the PDGui workbook."""
    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Crystal Raman"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    legends = [str(value) for value in header[2:] if value is not None]
    data = [row for row in rows if row[1] is not None]
    frequencies = np.asarray([row[1] for row in data], dtype=float)
    spectra = {
        legend: np.asarray([row[index + 2] for row in data], dtype=float)
        for index, legend in enumerate(legends)
    }
    return frequencies, spectra


def summarise(frequencies: np.ndarray, spectra: dict[str, np.ndarray]) -> list[dict]:
    """Extract peak and integrated-intensity comparisons in the polar band."""
    mask = (frequencies >= 450.0) & (frequencies <= 780.0)
    x = frequencies[mask]
    rows = []
    by_key = {}
    for legend, spectrum in spectra.items():
        match = LEGEND_RE.fullmatch(legend)
        if match is None:
            raise ValueError(f"Unexpected scenario legend: {legend!r}")
        y = spectrum[mask]
        peak_index = int(np.argmax(y))
        row = {
            "angle_external_deg": float(match.group("angle")),
            "nac_mode": match.group("nac"),
            "channel": match.group("channel"),
            "peak_frequency_cm-1": float(x[peak_index]),
            "peak_intensity": float(y[peak_index]),
            "integrated_intensity_450_780": float(np.trapezoid(y, x)),
        }
        rows.append(row)
        by_key[(row["angle_external_deg"], row["channel"], row["nac_mode"])] = row

    for row in rows:
        geometry = by_key[(row["angle_external_deg"], row["channel"], "geometry")]
        denominator = geometry["integrated_intensity_450_780"]
        row["integrated_ratio_to_geometry"] = (
            row["integrated_intensity_450_780"] / denominator
            if denominator != 0.0
            else float("nan")
        )
    return rows


def write_csv(rows: list[dict], output: Path) -> None:
    """Write a machine-readable summary."""
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_markdown(rows: list[dict], output: Path, workbook: Path) -> None:
    """Write a compact human-readable comparison table."""
    lines = [
        "# GaN modal-pairs results",
        "",
        f"Source workbook: `{workbook.as_posix()}`",
        "",
        "The table reports the strongest feature and integrated spectrum in the ",
        "450–780 cm⁻¹ polar-mode interval. Ratios are relative to `geometry` ",
        "for the same angle and optical channel.",
        "",
        "| external angle | NAC | channel | peak / cm⁻¹ | integrated intensity | ratio to geometry |",
        "|---:|---|---|---:|---:|---:|",
    ]
    for row in rows:
        lines.append(
            f"| {row['angle_external_deg']:g}° | {row['nac_mode']} | "
            f"{row['channel']} | {row['peak_frequency_cm-1']:.1f} | "
            f"{row['integrated_intensity_450_780']:.6g} | "
            f"{row['integrated_ratio_to_geometry']:.6g} |"
        )
    lines.extend(
        [
            "",
            "Interpretation must respect the model boundary: PDielec applies the ",
            "directional q→0 NAC correction. It tests optical-pair momentum routing, ",
            "but it does not calculate the finite-|q| phonon-polariton dispersion and ",
            "therefore is not expected to reproduce the experimental 485 cm⁻¹ peak.",
            "",
        ]
    )
    output.write_text("\n".join(lines))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook", nargs="?", type=Path, default=Path("Data/gan_modal_pairs.xlsx")
    )
    parser.add_argument("--csv", type=Path, default=Path("Data/gan_modal_pairs_summary.csv"))
    parser.add_argument("--markdown", type=Path, default=Path("RESULTS.md"))
    args = parser.parse_args()

    frequencies, spectra = load_spectra(args.workbook)
    rows = summarise(frequencies, spectra)
    write_csv(rows, args.csv)
    write_markdown(rows, args.markdown, args.workbook)
    print(f"Wrote {args.csv} and {args.markdown}")


if __name__ == "__main__":
    main()
